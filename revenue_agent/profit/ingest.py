"""File ingestion for investigations: CSV, XLSX (every worksheet) and text-based PDFs.

Pipeline: read file -> tables (header + rows) -> detect record type from headers -> map columns -> validate and
normalise rows into canonical records. Pure functions; persistence lives in store.py. Extracted values are
untrusted until validated; nothing is guessed. Uncertain detections or mappings are surfaced for the owner to
confirm instead of being applied silently.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date

from ..ledger.csv_import import (CURRENCY_EXPONENT, _norm_header, decode, parse_amount, parse_bool, parse_currency,
                                 parse_date)

MAX_BYTES = 10_000_000
MAX_ROWS = 20_000
MAX_CELL = 500
SUPPORTED = {".csv": "csv", ".txt": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".pdf": "pdf"}

# Canonical record types and their fields. `required` must be mapped; aliases are matched after normalisation.
TYPES: dict[str, dict[str, dict]] = {
    "sales": {
        "order_id": {"required": True, "aliases": ["order_id", "order", "order no", "order number", "order #",
                                                   "invoice", "invoice no", "رقم الطلب", "الطلب", "رقم الأوردر"]},
        "date": {"required": True, "aliases": ["date", "order date", "created at", "created_at", "sale date",
                                               "تاريخ", "التاريخ", "تاريخ الطلب"]},
        "product": {"aliases": ["product", "sku", "item", "product name", "item name", "lineitem name",
                                "lineitem sku", "المنتج", "الصنف", "اسم المنتج", "كود المنتج"]},
        "quantity": {"aliases": ["quantity", "qty", "units", "lineitem quantity", "الكمية", "العدد"]},
        "unit_price": {"aliases": ["unit price", "price", "lineitem price", "سعر الوحدة", "السعر"]},
        "revenue": {"aliases": ["revenue", "sales", "line total", "subtotal", "gross sales", "total", "amount",
                                "الإجمالي", "الاجمالي", "المبيعات", "قيمة الطلب"]},
        "discount": {"aliases": ["discount", "discount amount", "discounts", "lineitem discount", "خصم", "الخصم"]},
        "channel": {"aliases": ["channel", "sales channel", "source", "platform", "القناة", "منصة البيع"]},
        "status": {"aliases": ["status", "order status", "financial status", "الحالة", "حالة الطلب"]},
        "shipping_charged": {"aliases": ["shipping charged", "shipping paid", "delivery fee charged", "shipping fee",
                                         "shipping income", "رسوم التوصيل", "الشحن المدفوع", "مصاريف الشحن على العميل"]},
        "currency": {"aliases": ["currency", "العملة"]},
    },
    "product_costs": {
        "product": {"required": True, "aliases": ["product", "sku", "item", "product name", "المنتج", "الصنف",
                                                  "كود المنتج"]},
        "unit_cost": {"required": True, "aliases": ["unit cost", "cost", "cost per unit", "cogs", "purchase cost",
                                                    "manufacturing cost", "تكلفة الوحدة", "التكلفة", "سعر الشراء"]},
        "currency": {"aliases": ["currency", "العملة"]},
    },
    "expenses": {
        "date": {"required": True, "aliases": ["date", "invoice date", "bill date", "expense date", "التاريخ",
                                               "تاريخ الفاتورة"]},
        "amount": {"required": True, "aliases": ["amount", "total", "cost", "value", "invoice total", "delivery fee",
                                                 "shipping fee", "courier fee", "fee", "charge", "المبلغ", "القيمة",
                                                 "الإجمالي", "الاجمالي", "رسوم التوصيل", "رسوم الشحن"]},
        "category": {"aliases": ["category", "type", "expense type", "account", "البند", "النوع", "الفئة"]},
        "description": {"aliases": ["description", "details", "item", "memo", "notes", "الوصف", "البيان", "التفاصيل"]},
        "supplier": {"aliases": ["supplier", "vendor", "payee", "provider", "المورد", "الجهة"]},
        "quantity": {"aliases": ["quantity", "qty", "units", "الكمية"]},
        "invoice_id": {"aliases": ["invoice", "invoice id", "invoice no", "bill no", "reference", "رقم الفاتورة"]},
        "paid": {"aliases": ["paid", "is paid", "payment status", "status", "مدفوع", "حالة الدفع"]},
        "due_date": {"aliases": ["due date", "due", "تاريخ الاستحقاق"]},
        "order_id": {"aliases": ["order_id", "order", "order no", "order number", "order #", "order ref",
                                 "merchant reference", "رقم الطلب"]},
        "currency": {"aliases": ["currency", "العملة"]},
    },
    "returns": {
        "order_id": {"required": True, "aliases": ["order_id", "order", "order no", "order number", "رقم الطلب"]},
        "date": {"required": True, "aliases": ["date", "return date", "refund date", "تاريخ", "تاريخ المرتجع"]},
        "product": {"aliases": ["product", "sku", "item", "المنتج", "الصنف"]},
        "quantity": {"aliases": ["quantity", "qty", "units", "الكمية"]},
        "refund_amount": {"aliases": ["refund", "refund amount", "amount refunded", "amount", "قيمة المرتجع",
                                      "المبلغ المسترد", "المبلغ"]},
        "reason": {"aliases": ["reason", "return reason", "سبب", "سبب المرتجع"]},
        "currency": {"aliases": ["currency", "العملة"]},
    },
}
TYPE_HINTS = {  # header words that strongly suggest one type over another
    "returns": ["return", "refund", "مرتجع", "مسترد"],
    "product_costs": ["unit cost", "cost per unit", "cogs", "تكلفة الوحدة"],
    "expenses": ["supplier", "vendor", "category", "expense", "courier", "awb", "tracking", "delivery fee",
                 "shipping fee", "fee", "commission", "المورد", "البند", "شحن", "عمولة"],
    "sales": ["order", "qty", "quantity", "طلب"],
}

# Expense categories. Variable = scales with orders; fixed = does not.
CATEGORIES = ("cogs_materials", "packaging", "shipping", "payment_fees", "marketplace_fees", "marketing", "rent",
              "payroll", "subscriptions", "other")
VARIABLE = ("packaging", "shipping", "payment_fees", "marketplace_fees")
FIXED = ("rent", "payroll", "subscriptions", "other")
_CATEGORY_WORDS: list[tuple[str, tuple[str, ...]]] = [
    ("packaging", ("packag", "box", "boxes", "carton", "bag", "mailer", "label", "wrap", "تغليف", "كرتون", "علب",
                   "علبة", "شنط", "اكياس", "أكياس", "ليبل")),
    ("shipping", ("shipping", "delivery", "courier", "freight", "postage", "bosta", "aramex", "mylerz", "dhl",
                  "fedex", "شحن", "توصيل", "مندوب", "بوسطة", "أرامكس")),
    ("payment_fees", ("payment fee", "processing", "paymob", "fawry", "stripe", "paypal", "gateway", "card fee",
                      "رسوم دفع", "فوري", "بوابة دفع")),
    ("marketplace_fees", ("commission", "marketplace", "amazon", "noon", "jumia", "etsy", "عمولة", "نون", "جوميا")),
    ("marketing", ("marketing", "ads", "advert", "facebook", "instagram", "meta", "google ads", "tiktok",
                   "influencer", "اعلان", "إعلان", "اعلانات", "تسويق", "سوشيال")),
    ("rent", ("rent", "lease", "إيجار", "ايجار")),
    ("payroll", ("salary", "salaries", "payroll", "wage", "رواتب", "مرتبات", "أجور", "اجور")),
    ("subscriptions", ("subscription", "software", "saas", "shopify", "hosting", "domain", "اشتراك", "استضافة")),
    ("cogs_materials", ("fabric", "material", "raw", "inventory purchase", "stock purchase", "قماش", "خامات",
                        "خامة", "مواد", "بضاعة")),
]
_STATUS_CANCELLED = {"cancelled", "canceled", "void", "voided", "failed", "ملغي", "ملغى", "ملغاة", "اتلغى"}
_STATUS_RETURNED = {"returned", "refunded", "مرتجع", "مسترد", "اترجع"}


@dataclass
class Issue:
    sheet: str
    row: int
    field: str
    code: str
    message: str
    severity: str = "error"  # error (row rejected) | warning (row kept, flagged)

    def as_dict(self) -> dict:
        return self.__dict__.copy()


@dataclass
class TableResult:
    sheet: str
    headers: list[str]
    detected_type: str | None
    type_scores: dict[str, float]
    mapping: dict[str, str]
    missing_required: list[str]
    needs_confirmation: list[str] = field(default_factory=list)
    records: list[dict] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    rows_total: int = 0

    def summary(self) -> dict:
        return {"sheet": self.sheet, "headers": self.headers, "detected_type": self.detected_type,
                "type_scores": self.type_scores, "mapping": self.mapping, "missing_required": self.missing_required,
                "needs_confirmation": self.needs_confirmation, "rows_total": self.rows_total,
                "rows_valid": len(self.records), "rows_rejected": len({(i.row) for i in self.issues
                                                                        if i.severity == "error"}),
                "issues": [i.as_dict() for i in self.issues[:300]], "issue_count": len(self.issues),
                "periods": sorted({r["date"][:7] for r in self.records if r.get("date")}),
                "currencies": sorted({r["currency"] for r in self.records if r.get("currency")})}


@dataclass
class FileResult:
    filename: str
    kind: str | None
    sha256: str
    tables: list[TableResult] = field(default_factory=list)
    fatal: str | None = None
    text_preview: str | None = None

    def summary(self) -> dict:
        return {"filename": self.filename, "kind": self.kind, "sha256": self.sha256, "fatal": self.fatal,
                "tables": [t.summary() for t in self.tables], "text_preview": self.text_preview}


# ------------------------------------------------------------------------------------------ readers

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, date):
        return v.isoformat()[:10]
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = unicodedata.normalize("NFKC", str(v)).strip()
    return s[:MAX_CELL]


def read_tables(filename: str, raw: bytes) -> tuple[str | None, list[tuple[str, list[list[str]]]], str | None]:
    """Returns (kind, [(sheet_name, rows)], text_preview). Raises ValueError with a user-facing reason."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    kind = SUPPORTED.get(ext)
    if kind is None:
        raise ValueError(f"unsupported file type '{ext or 'none'}': upload CSV, XLSX or a text-based PDF")
    if len(raw) > MAX_BYTES:
        raise ValueError(f"file larger than {MAX_BYTES // 1_000_000} MB")
    if not raw.strip():
        raise ValueError("file is empty")
    if kind == "csv":
        text = decode(raw)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return kind, [("csv", [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), dialect)])], None
    if kind == "xlsx":
        if not raw.startswith(b"PK"):
            raise ValueError("this is not a valid .xlsx workbook (old .xls files are not supported: save as .xlsx)")
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for corrupt files
            raise ValueError(f"could not open the workbook ({type(exc).__name__})") from exc
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS + 1:
                    break
                rows.append([_cell(c) for c in r])
            sheets.append((ws.title[:60], rows))
        wb.close()
        return kind, sheets, None
    # pdf: text extraction only (no OCR). Tables are recovered when lines are delimiter-separated.
    import pypdf
    try:
        reader = pypdf.PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("the PDF is password-protected")
        text = "\n".join((p.extract_text() or "") for p in reader.pages[:50])
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not read the PDF ({type(exc).__name__})") from exc
    if len(text.strip()) < 20:
        raise ValueError("no text found in the PDF: it is probably a scanned image. OCR is not supported yet; "
                         "type the key numbers in manually or upload the CSV/Excel version")
    rows = _pdf_rows(text)
    return kind, ([("pdf", rows)] if rows else []), text[:1500]


def _pdf_rows(text: str) -> list[list[str]]:
    """Recover a table from PDF text when lines use a consistent delimiter (|, ;, tab, or 2+ spaces)."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for splitter in (r"\s*\|\s*", r"\s*;\s*", r"\t+", r"\s{2,}"):
        rows = [re.split(splitter, l) for l in lines]
        widths = [len(r) for r in rows]
        best = max(set(widths), key=widths.count) if widths else 0
        table = [r for r in rows if len(r) == best]
        if best >= 3 and len(table) >= 2:
            return [[_cell(c) for c in r] for r in table]
    return []


# ------------------------------------------------------------------------------------------ detection + mapping

def _find_header(rows: list[list[str]]) -> int:
    """First row (within the top 10) that looks like a header: mostly non-numeric, ≥2 non-empty cells."""
    for i, r in enumerate(rows[:10]):
        cells = [c for c in r if c]
        if len(cells) >= 2 and sum(not re.fullmatch(r"[\d.,\-/ ]+", c) for c in cells) >= max(2, len(cells) * 0.6):
            return i
    return 0


def map_columns(headers: list[str], rtype: str) -> dict[str, str]:
    mapping, used = {}, set()
    normed = {h: _norm_header(h) for h in headers if h}
    for fname, spec in TYPES[rtype].items():
        aliases = [_norm_header(a) for a in spec["aliases"]]
        for h, n in normed.items():  # exact alias first
            if h not in used and n in aliases:
                mapping[fname] = h
                used.add(h)
                break
    return mapping


def detect_type(headers: list[str]) -> tuple[str | None, dict[str, float], list[str]]:
    """Score each record type by how many of its fields map and whether required ones are present."""
    scores: dict[str, float] = {}
    low = " ".join(_norm_header(h) for h in headers)
    for rtype, spec in TYPES.items():
        m = map_columns(headers, rtype)
        req = [f for f, s in spec.items() if s.get("required")]
        if not all(f in m for f in req) or (rtype == "sales" and not ({"revenue", "unit_price"} & set(m))):
            scores[rtype] = 0.0  # a sales table must say what the customer paid
            continue
        hint = sum(w in low for w in TYPE_HINTS[rtype])
        scores[rtype] = round(len(m) / len(spec) + 0.15 * hint, 3)
    ranked = sorted(scores.items(), key=lambda kv: -kv[1])
    notes = []
    if not ranked or ranked[0][1] == 0:
        return None, scores, ["could not tell what this table contains: please choose the type and map columns"]
    if len(ranked) > 1 and ranked[1][1] > 0 and ranked[0][1] - ranked[1][1] < 0.15:
        notes.append(f"this table could be '{ranked[0][0]}' or '{ranked[1][0]}': please confirm the type")
    return ranked[0][0], scores, notes


def categorize(category: str, description: str, supplier: str) -> tuple[str, bool]:
    """Expense category from an explicit column if recognisable, else inferred from text (flagged as inferred)."""
    for text, inferred in ((category, False), (f"{description} {supplier}", True)):
        t = (text or "").lower()
        if not t.strip():
            continue
        if not inferred and t.strip() in CATEGORIES:
            return t.strip(), False
        for cat, words in _CATEGORY_WORDS:
            if any(w in t for w in words):
                return cat, inferred
    return "other", True


# ------------------------------------------------------------------------------------------ normalisation

def _num(raw: str, exp: int, field_name: str, issues: list, sheet: str, row: int, allow_neg=False) -> int | None:
    if not raw:
        return None
    try:
        v = parse_amount(raw, exp)
    except ValueError as exc:
        issues.append(Issue(sheet, row, field_name, "invalid_number", f"{field_name}: {exc}"))
        return None
    if v < 0 and not allow_neg:
        issues.append(Issue(sheet, row, field_name, "negative", f"{field_name} is negative"))
        return None
    return v


def _qty(raw: str, issues: list, sheet: str, row: int) -> float | None:
    if not raw:
        return None
    try:
        q = float(raw.replace(",", "").translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")))
    except ValueError:
        issues.append(Issue(sheet, row, "quantity", "invalid_number", "quantity is not a number"))
        return None
    if q <= 0:
        issues.append(Issue(sheet, row, "quantity", "non_positive", "quantity must be positive"))
        return None
    return q


def normalise(sheet: str, headers: list[str], rows: list[list[str]], rtype: str, mapping: dict[str, str], *,
              default_currency: str, date_format: str, first_row: int) -> TableResult:
    spec = TYPES[rtype]
    missing = [f for f, s in spec.items() if s.get("required") and f not in mapping]
    res = TableResult(sheet, headers, rtype, {}, mapping, missing)
    if missing:
        return res
    if rtype == "sales" and not ({"revenue", "unit_price"} & set(mapping)):
        res.missing_required.append("revenue|unit_price")
        return res
    idx = {f: headers.index(h) for f, h in mapping.items() if h in headers}
    seen: dict[tuple, int] = {}
    for n, cells in enumerate(rows, start=first_row):
        if not any(c for c in cells):
            continue
        res.rows_total += 1
        if res.rows_total > MAX_ROWS:
            res.issues.append(Issue(sheet, n, "-", "too_many_rows", f"more than {MAX_ROWS} rows; the rest was ignored"))
            break
        get = lambda f: (cells[idx[f]] if f in idx and idx[f] < len(cells) else "").strip()  # noqa: E731
        errs: list[Issue] = []
        cur_raw = get("currency")
        cur = parse_currency(cur_raw) if cur_raw else default_currency
        if cur is None:
            errs.append(Issue(sheet, n, "currency", "unknown_currency", f"unknown currency '{cur_raw}'"))
            cur = default_currency
        exp = CURRENCY_EXPONENT.get(cur, 2)
        rec: dict = {"type": rtype, "currency": cur, "sheet": sheet, "row": n, "flags": []}
        d_raw = get("date") if "date" in spec else ""
        if "date" in spec:
            if not d_raw:
                errs.append(Issue(sheet, n, "date", "required", "date is empty"))
            else:
                try:
                    d, amb = parse_date(d_raw, date_format)
                    rec["date"] = d.isoformat()
                    if amb:
                        res.issues.append(Issue(sheet, n, "date", "ambiguous_date",
                                                f"{d_raw} read as {d.isoformat()} (day-first)", "warning"))
                except ValueError as exc:
                    errs.append(Issue(sheet, n, "date", "invalid_date", f"date: {exc}"))
        if rtype == "sales":
            rec["order_id"] = get("order_id")[:64]
            if not rec["order_id"]:
                errs.append(Issue(sheet, n, "order_id", "required", "order ID is empty"))
            rec["product"] = get("product")[:120] or None
            q = _qty(get("quantity"), errs, sheet, n)
            rec["quantity"] = q if q is not None else 1.0
            if q is None and get("quantity") == "":
                rec["flags"].append("quantity_assumed_1")
            price = _num(get("unit_price"), exp, "unit_price", errs, sheet, n)
            rev = _num(get("revenue"), exp, "revenue", errs, sheet, n)
            if rev is None and price is not None:
                rev = round(price * rec["quantity"])
            if rev is None:
                errs.append(Issue(sheet, n, "revenue", "required", "no revenue or unit price"))
            elif price is not None and abs(rev - price * rec["quantity"]) > max(100, 0.02 * rev) and "revenue" in idx:
                res.issues.append(Issue(sheet, n, "revenue", "price_qty_mismatch",
                                        "line total ≠ price × quantity (line total used)", "warning"))
            rec["gross_minor"] = rev
            disc = _num(get("discount"), exp, "discount", errs, sheet, n, allow_neg=True)
            rec["discount_minor"] = abs(disc) if disc is not None else 0
            if rev is not None and rec["discount_minor"] > rev:
                errs.append(Issue(sheet, n, "discount", "discount_exceeds_revenue", "discount larger than the line"))
            rec["channel"] = get("channel")[:60] or None
            rec["shipping_charged_minor"] = _num(get("shipping_charged"), exp, "shipping_charged", errs, sheet, n) or 0
            st = get("status").lower()
            rec["status"] = ("cancelled" if st in _STATUS_CANCELLED else "returned" if st in _STATUS_RETURNED
                             else "completed")
            key = (rec["order_id"], rec.get("product"), rec.get("date"), rec["gross_minor"])
        elif rtype == "product_costs":
            rec["product"] = get("product")[:120]
            rec["unit_cost_minor"] = _num(get("unit_cost"), exp, "unit_cost", errs, sheet, n)
            if rec["unit_cost_minor"] is None and not any(e.field == "unit_cost" for e in errs):
                errs.append(Issue(sheet, n, "unit_cost", "required", "unit cost is empty"))
            key = (rec["product"].lower(),)
        elif rtype == "expenses":
            rec["amount_minor"] = _num(get("amount"), exp, "amount", errs, sheet, n)
            if rec["amount_minor"] is None and not any(e.field == "amount" for e in errs):
                errs.append(Issue(sheet, n, "amount", "required", "amount is empty"))
            rec["description"] = get("description")[:200] or None
            rec["supplier"] = get("supplier")[:120] or None
            cat, inferred = categorize(get("category"), rec["description"] or "", rec["supplier"] or "")
            rec["category"] = cat
            if inferred:
                rec["flags"].append("category_inferred")
            rec["quantity"] = _qty(get("quantity"), errs, sheet, n)
            rec["invoice_id"] = get("invoice_id")[:64] or None
            rec["order_id"] = get("order_id")[:64] or None  # links a cost to one order (courier/gateway statements)
            paid = parse_bool(get("paid")) if get("paid") else None
            if get("paid") and paid is None:
                low = get("paid").lower()
                paid = True if low in ("paid", "settled", "مدفوع", "مدفوعة") else False if low in (
                    "unpaid", "due", "open", "غير مدفوع", "مستحق") else None
            rec["paid"] = paid
            due = get("due_date")
            if due:
                try:
                    rec["due_date"] = parse_date(due, date_format)[0].isoformat()
                except ValueError:
                    res.issues.append(Issue(sheet, n, "due_date", "invalid_date", "due date unreadable", "warning"))
            key = (rec.get("invoice_id") or "", rec.get("supplier") or "", rec.get("date"), rec["amount_minor"],
                   rec.get("description") or "", rec.get("order_id") or "", rec.get("category"))
        else:  # returns
            rec["order_id"] = get("order_id")[:64]
            if not rec["order_id"]:
                errs.append(Issue(sheet, n, "order_id", "required", "order ID is empty"))
            rec["product"] = get("product")[:120] or None
            rec["quantity"] = _qty(get("quantity"), errs, sheet, n) or 1.0
            rec["refund_minor"] = _num(get("refund_amount"), exp, "refund_amount", errs, sheet, n) or 0
            rec["reason"] = get("reason")[:200] or None
            key = (rec["order_id"], rec.get("product"), rec.get("date"), rec["refund_minor"])
        if errs:
            res.issues.extend(errs)
            continue
        if key in seen:
            res.issues.append(Issue(sheet, n, "-", "duplicate_record", f"same as row {seen[key]}; skipped", "warning"))
            continue
        seen[key] = n
        res.records.append(rec)
    if rtype == "expenses":
        inferred = sum("category_inferred" in r["flags"] for r in res.records)
        if inferred:
            res.needs_confirmation.append(f"{inferred} expense rows had their category inferred from the "
                                          "description/supplier: please review the categories")
    return res


def parse_file(filename: str, raw: bytes, *, default_currency: str = "EGP", date_format: str = "auto",
               overrides: dict[str, dict] | None = None) -> FileResult:
    """overrides: {sheet: {"type": ..., "mapping": {...}}} confirmed by the owner (or reused from memory)."""
    fr = FileResult(filename[:200], None, hashlib.sha256(raw).hexdigest())
    try:
        fr.kind, sheets, fr.text_preview = read_tables(filename, raw)
    except ValueError as exc:
        fr.fatal = str(exc)
        return fr
    if not sheets:
        fr.fatal = ("no table found in this PDF. Text-based invoice PDFs are read when the lines form a table; "
                    "otherwise type the totals in manually")
        return fr
    for sheet, rows in sheets:
        if not any(any(c for c in r) for r in rows):
            continue
        h = _find_header(rows)
        headers = [c or f"column_{i + 1}" for i, c in enumerate(rows[h])]
        if len(set(headers)) != len(headers):
            fr.tables.append(TableResult(sheet, headers, None, {}, {}, ["duplicate column names"]))
            continue
        ov = (overrides or {}).get(sheet) or {}
        rtype, scores, notes = detect_type(headers)
        if ov.get("type") in TYPES:
            rtype, notes = ov["type"], []
        if rtype is None:
            fr.tables.append(TableResult(sheet, headers, None, scores, {}, [], notes))
            continue
        mapping = ov.get("mapping") or map_columns(headers, rtype)
        mapping = {k: v for k, v in mapping.items() if k in TYPES[rtype] and v in headers}
        t = normalise(sheet, headers, rows[h + 1:], rtype, mapping, default_currency=default_currency,
                      date_format=date_format, first_row=h + 2)
        t.type_scores, t.needs_confirmation = scores, notes + t.needs_confirmation
        fr.tables.append(t)
    if not fr.tables:
        fr.fatal = "the file has no data rows"
    return fr


def header_signature(headers: list[str]) -> str:
    """Stable fingerprint of a report layout, used to reuse an owner-confirmed mapping for the same format."""
    return hashlib.sha256("|".join(sorted(_norm_header(h) for h in headers if h)).encode()).hexdigest()[:16]
